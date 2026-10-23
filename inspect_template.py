from pptx import Presentation

prs = Presentation('MQI_Natl_Template.pptx')
for idx, slide in enumerate(prs.slides, 1):
    print(f"\n=================== SLIDE {idx} ===================")
    for s in slide.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip().replace('\n', ' // ')
            if t:
                print(f"  [TEXT] {s.name:<25} | {t[:100]}")
        if s.has_chart:
            c = s.chart
            title = c.chart_title.text_frame.text if c.has_title else "No Title"
            series = [series.name for series in c.series]
            cats = [str(cat) for cat in c.plots[0].categories] if c.plots and hasattr(c.plots[0], 'categories') else []
            print(f"  [CHART] {s.name:<24} | Title: {title} | Series: {series} | Categories: {cats[:6]}")
