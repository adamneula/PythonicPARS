"""
presentationTaxible.py
----------------------
Generates a PowerPoint presentation for taxable Corporate / Treasury / Agency portfolios.
Modeled after Genter's institutional standard:
    PROPOSAL PARS INSTITUTIONAL V4.5 - TAXABLE 2026

Features & Mappings:
  - Slide 1: Strategy updated to "9100 - Taxable Quality Intermediate".
  - Slide 2: Preserved as placeholder / credit risk assessment.
  - Slide 3: Structure Analysis (Corporate vs Treasury) & Industry Breakdown
             (using industry_subgroup/muni_purpose to match PDF industry slices).
  - Slide 4: Structure Recommendation:
             Current Portfolio vs Recommended Portfolio (Treasury 50% / Corporate 50%).
  - Slide 5: Stated Maturity Distribution & Callability.
  - Slide 6: Maturity Recommendation:
             Current Adjusted Maturity vs Recommended Portfolio
             (1-3 Yr: 17.5%, 3-5 Yr: 30.0%, 5-7 Yr: 40.0%, 7-10 Yr: 10.0%, 15+ Yr: 2.5%).
  - Slide 7: Coupon Rate Analysis using PDF 4-tier buckets (0%-3%, 3%-4%, 4%-5%, 5%+).
  - Slide 8: Credit Quality Recommendation:
             Current Portfolio vs Recommended Portfolio (Treasury 50.0%, BBB 30.0%, A 20.0%).
  - Slide 9: Marks the Municipal State Breakdown slide as hidden (show="0"), so it is omitted
             from the presentation flow without XML file corruption or warnings.
  - Slide 10: Characteristics Comparison updated with Genter's Taxable Quality Intermediate
              benchmarks (Coupon 4.42%, YTW 4.61%, Eff Mat 6.10 Yrs, Eff Dur 3.98 Yrs, Conv 0.23)
              and sets the federal tax footnote to 0.0%.
  - Slide 11+: Current Portfolio Investment Detail by Maturity table populated directly onto
               the template's HoldingsDetail slide (adding additional slides only if needed).
  - Safe Saving: Automatically falls back to timestamped file if PowerPoint holds a lock.

Usage:
    python presentationTaxible.py
"""

import os
import glob
import datetime
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from mappings import (
    CONSOLIDATED_RATINGS_ORDER,
    MOODYS_TO_SP,
    RATING_SCORES,
    SCORE_TO_RATING,
)


# ── Helper functions ──────────────────────────────────────────────────────────

def get_shape(slide, shape_name):
    """Searches ONLY the current slide for a shape by name."""
    for shape in slide.shapes:
        if shape.name == shape_name:
            return shape
    return None


def update_text_preserve_format(shape, new_text):
    """
    Overwrites text in a shape while preserving the exact font, size,
    color, and bold/italic styling set in PowerPoint.
    """
    if not shape.has_text_frame:
        return

    p = shape.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(new_text)
        for run in p.runs[1:]:
            run.text = ""
    else:
        shape.text = str(new_text)


def synthesize_rating(row):
    """Resolves composite rating with S&P, Moody's, and Fitch fallbacks."""
    comp = str(row.get('composite_rating', 'NR')).strip()
    if comp and comp != 'NR' and comp != 'nan':
        return comp

    sp = str(row.get('sp_rating', 'NR')).strip()
    if sp and sp != 'NR' and sp != 'nan':
        return sp

    mdy = str(row.get('moodys_rating', 'NR')).strip()
    if mdy and mdy != 'NR' and mdy != 'nan':
        return MOODYS_TO_SP.get(mdy, mdy)

    fitch = str(row.get('fitch_rating', 'NR')).strip()
    if fitch and fitch != 'NR' and fitch != 'nan':
        return fitch

    return 'NR'


def find_latest_portfolio_file(explicit_path=None):
    """Finds the most appropriate portfolio output file to process."""
    if explicit_path and os.path.exists(explicit_path):
        return explicit_path

    # Priority 1: Newly generated taxable files
    taxable_files = sorted(glob.glob("Portfolio_Output_Taxable_*.xlsx"))
    if taxable_files:
        return taxable_files[-1]

    # Priority 2: Hardcoded run from today
    default_today = "Portfolio_Output_20261020_093547.xlsx"
    if os.path.exists(default_today):
        return default_today

    # Priority 3: Any Portfolio_Output_*.xlsx
    any_files = sorted(glob.glob("Portfolio_Output_*.xlsx"))
    if any_files:
        return any_files[-1]

    raise FileNotFoundError("Could not find any Portfolio_Output Excel file to load.")


# ── Presentation Generator ───────────────────────────────────────────────────

def generate_taxable_presentation(
    advisor_name: str = "Jeff Kuchinsky, CFP, MBA\nJMK Wealth Advisors",
    client_name: str = "Prospective Client",
    excel_path: str = None,
    template_path: str = "MQI_Natl_Template.pptx",
):
    """
    Builds the taxable presentation from Excel data and PPTX template,
    matching the Genter V4.5 Taxable Institutional standard.
    """
    # 1. Resolve & Load Data
    resolved_file = find_latest_portfolio_file(excel_path)
    print(f"Reading portfolio data from: {resolved_file}")
    df = pd.read_excel(resolved_file)

    total_value = df['market_value'].sum()

    # 2. Derive Taxable Classification Columns
    if 'security_type' in df.columns and (df['security_type'] != 'Unknown').any():
        df['clean_security_type'] = df['security_type']
    else:
        df['clean_security_type'] = 'Corporate'

    # Industry Classification: for Slide 3 Right Chart (matches PDF "Income Source")
    if 'industry_group' in df.columns and (df['industry_group'] != 'Unknown').any():
        df['clean_industry'] = df['industry_group']
    elif 'industry_subgroup' in df.columns and (df['industry_subgroup'] != 'Unknown').any():
        df['clean_industry'] = df['industry_subgroup']
    elif 'muni_purpose' in df.columns and (df['muni_purpose'] != 'Unknown').any():
        df['clean_industry'] = df['muni_purpose']
    else:
        df['clean_industry'] = 'Corporate'

    # 3. Load PowerPoint Template
    print(f"Loading PowerPoint template: {template_path}")
    prs = Presentation(template_path)

    # ========================================================
    # SLIDE 1: Title & Overview
    # ========================================================
    print("Populating Slide 1...")
    slide1 = prs.slides[0]

    strat_box = get_shape(slide1, "Strategy + Portfolio Amt")
    if strat_box:
        update_text_preserve_format(strat_box, "9100 - Taxable Quality Intermediate in the amount of:")

    client_box = get_shape(slide1, "ClientName")
    if client_box:
        update_text_preserve_format(client_box, f"Prepared For: {client_name}")

    advisor_box = get_shape(slide1, "AdvisorName")
    if advisor_box:
        update_text_preserve_format(advisor_box, f"Per the request of:\n{advisor_name}")

    money_box = get_shape(slide1, "ValueBox")
    if money_box:
        update_text_preserve_format(money_box, f"${total_value:,.0f}")

    # ========================================================
    # SLIDE 2: Credit Risk Assessment (Placeholder)
    # ========================================================
    print("Populating Slide 2...")
    # Slide 2 preserved as-is

    # ========================================================
    # SLIDE 3: Structure Analysis – Current Portfolio
    # ========================================================
    print("Populating Slide 3 (Structure & Industry)...")
    slide3 = prs.slides[2]

    # Chart 1: Structure Analysis (Corporate vs Treasury)
    struct_counts = df.groupby('clean_security_type')['market_value'].sum()
    struct_chart_data = CategoryChartData()
    struct_chart_data.categories = [str(lbl).title() for lbl in struct_counts.index]
    struct_chart_data.add_series('Structure', struct_counts.values.tolist())

    chart_s3_1 = get_shape(slide3, "ClientStructureAnalysis")
    if chart_s3_1 and chart_s3_1.has_chart:
        chart_s3_1.chart.replace_data(struct_chart_data)

    # Chart 2: Industry Breakdown (mapped to ClientIncomeSource shape)
    ind_counts = df.groupby('clean_industry')['market_value'].sum().sort_values(ascending=False)
    ind_chart_data = CategoryChartData()
    ind_chart_data.categories = [str(c).title() for c in ind_counts.index.tolist()]
    ind_chart_data.add_series('Industries', ind_counts.values.tolist())

    chart_s3_2 = get_shape(slide3, "ClientIncomeSource")
    if chart_s3_2 and chart_s3_2.has_chart:
        chart_s3_2.chart.replace_data(ind_chart_data)

    # ========================================================
    # SLIDE 4: Structure Recommendation
    # ========================================================
    print("Populating Slide 4 (Structure Recommendation)...")
    slide4 = prs.slides[3]

    # Left Chart: Current Portfolio
    chart_s4_curr = get_shape(slide4, "ClientStructure")
    if chart_s4_curr and chart_s4_curr.has_chart:
        chart_s4_curr.chart.replace_data(struct_chart_data)

    # Right Chart: Recommended Portfolio -> Treasury 50%, Corporate 50% (from PDF Page 4)
    chart_s4_rec = get_shape(slide4, "GenterStructure")
    if chart_s4_rec and chart_s4_rec.has_chart:
        rec_struct_data = CategoryChartData()
        rec_struct_data.categories = ['Treasury', 'Corporate']
        rec_struct_data.add_series('Recommended Portfolio', [50.0, 50.0])
        chart_s4_rec.chart.replace_data(rec_struct_data)

    # ========================================================
    # SLIDE 5: Duration & Maturity Analysis – Current Portfolio
    # ========================================================
    print("Populating Slide 5 (Maturity & Callability)...")
    slide5 = prs.slides[4]

    # Stated Maturity
    df['maturity_date'] = pd.to_datetime(df['maturity_date'], errors='coerce')
    today = pd.Timestamp.today()
    df['years_to_maturity'] = (df['maturity_date'] - today).dt.days / 365.25

    def bucket_maturity(years):
        if pd.isna(years): return "Unknown"
        if years < 1:   return "0-1"
        elif years < 3: return "1-3"
        elif years < 5: return "3-5"
        elif years < 7: return "5-7"
        elif years < 10: return "7-10"
        elif years < 15: return "10-15"
        else:           return "15+"

    df['maturity_bucket'] = df['years_to_maturity'].apply(bucket_maturity)
    mat_counts = df.groupby('maturity_bucket')['market_value'].sum()

    ordered_mat = ["0-1", "1-3", "3-5", "5-7", "7-10", "10-15", "15+"]
    ordered_labels_mat = [r for r in ordered_mat if r in mat_counts.index]
    mat_counts = mat_counts.reindex(ordered_labels_mat)

    chart_data_mat = CategoryChartData()
    chart_data_mat.categories = mat_counts.index.tolist()
    chart_data_mat.add_series('Stated Maturity', mat_counts.values.tolist())

    chart_s5_mat = get_shape(slide5, "ClientMaturityProfile")
    if chart_s5_mat and chart_s5_mat.has_chart:
        chart_s5_mat.chart.replace_data(chart_data_mat)

    # Callability
    def clean_callable(val):
        val_str = str(val).strip().upper()
        if val_str in ["TRUE", "Y", "YES", "1", "1.0"]:
            return "Callable"
        return "Non-Callable"

    df['clean_callable'] = df.get('is_callable', False).apply(clean_callable)
    call_counts = df.groupby('clean_callable')['market_value'].sum()

    chart_data_call = CategoryChartData()
    chart_data_call.categories = call_counts.index.tolist()
    chart_data_call.add_series('Callability', call_counts.values.tolist())

    chart_s5_call = get_shape(slide5, "ClientCallRate")
    if chart_s5_call and chart_s5_call.has_chart:
        chart_s5_call.chart.replace_data(chart_data_call)

    # ========================================================
    # SLIDE 6: Maturity Recommendation
    # ========================================================
    print("Populating Slide 6 (Maturity Recommendation)...")
    slide6 = prs.slides[5]

    # Left Chart: Current Adjusted Maturity
    df['effective_maturity_date'] = pd.to_datetime(df.get('effective_maturity_date'), errors='coerce')
    df['years_to_eff_maturity'] = (df['effective_maturity_date'] - today).dt.days / 365.25
    df['eff_maturity_bucket'] = df['years_to_eff_maturity'].apply(bucket_maturity)

    eff_mat_counts = df.groupby('eff_maturity_bucket')['market_value'].sum()
    ordered_labels_eff = [r for r in ordered_mat if r in eff_mat_counts.index]
    eff_mat_counts = eff_mat_counts.reindex(ordered_labels_eff)

    chart_data_eff = CategoryChartData()
    chart_data_eff.categories = eff_mat_counts.index.tolist()
    chart_data_eff.add_series('Current Adjusted Maturity', eff_mat_counts.values.tolist())

    chart_s6_eff = get_shape(slide6, "ClientOAM")
    if chart_s6_eff and chart_s6_eff.has_chart:
        chart_s6_eff.chart.replace_data(chart_data_eff)

    # Right Chart: Recommended Portfolio (from PDF Page 6)
    # 1-3 Years: 17.5%, 3-5 Years: 30.0%, 5-7 Years: 40.0%, 7-10 Years: 10.0%, 15+ Years: 2.5%
    chart_s6_genter = get_shape(slide6, "GenterOAM")
    if chart_s6_genter and chart_s6_genter.has_chart:
        genter_mat_data = CategoryChartData()
        genter_mat_data.categories = ['1-3 Years', '3-5 Years', '5-7 Years', '7-10 Years', '15+ Years']
        genter_mat_data.add_series('Recommended Portfolio', [17.5, 30.0, 40.0, 10.0, 2.5])
        chart_s6_genter.chart.replace_data(genter_mat_data)

    # ========================================================
    # SLIDE 7: Coupon Rate Analysis – Current Portfolio
    # ========================================================
    print("Populating Slide 7 (Coupon Rate)...")
    slide7 = prs.slides[6]

    df['coupon'] = pd.to_numeric(df.get('coupon'), errors='coerce')
    if df['coupon'].max() < 0.20 and df['coupon'].max() > 0:
        df['coupon_clean'] = df['coupon'] * 100
    else:
        df['coupon_clean'] = df['coupon']

    # 4-tier buckets matching PDF Page 7
    def bucket_coupon_taxable(rate):
        if pd.isna(rate): return "Unknown"
        if rate < 3.0:    return "0%-3%"
        elif rate < 4.0:  return "3%-4%"
        elif rate < 5.0:  return "4%-5%"
        else:             return "5%+"

    df['coupon_bucket'] = df['coupon_clean'].apply(bucket_coupon_taxable)
    cpn_counts = df.groupby('coupon_bucket')['market_value'].sum()

    ordered_coupons = ["0%-3%", "3%-4%", "4%-5%", "5%+"]
    ordered_labels_cpn = [r for r in ordered_coupons if r in cpn_counts.index]
    cpn_counts = cpn_counts.reindex(ordered_labels_cpn)

    chart_data_cpn = CategoryChartData()
    chart_data_cpn.categories = cpn_counts.index.tolist()
    chart_data_cpn.add_series('Coupon Rate', cpn_counts.values.tolist())

    chart_s7_cpn = get_shape(slide7, "ClientCouponRate")
    if chart_s7_cpn and chart_s7_cpn.has_chart:
        chart_s7_cpn.chart.replace_data(chart_data_cpn)

    # ========================================================
    # SLIDE 8: Credit Quality Recommendation
    # ========================================================
    print("Populating Slide 8 (Quality Recommendation)...")
    slide8 = prs.slides[7]

    # Left Chart: Current Portfolio
    df['syn_rating'] = df.apply(synthesize_rating, axis=1)
    df['clean_rating'] = df['syn_rating'].str.replace('+', '', regex=False).str.replace('-', '', regex=False)
    ratings_counts = df.groupby('clean_rating')['market_value'].sum()

    ordered_labels_rtg = [r for r in CONSOLIDATED_RATINGS_ORDER if r in ratings_counts.index]
    ratings_counts = ratings_counts.reindex(ordered_labels_rtg)

    chart_data_ratings = CategoryChartData()
    chart_data_ratings.categories = ratings_counts.index.tolist()
    chart_data_ratings.add_series('Credit Quality', ratings_counts.values.tolist())

    chart_s8_curr = get_shape(slide8, "ClientQualityBreakdown")
    if chart_s8_curr and chart_s8_curr.has_chart:
        chart_s8_curr.chart.replace_data(chart_data_ratings)

    # Right Chart: Recommended Portfolio (from PDF Page 8)
    # Treasury: 50.0%, BBB: 30.0%, A: 20.0%
    chart_s8_rec = get_shape(slide8, "GenterQualityBreakdown")
    if chart_s8_rec and chart_s8_rec.has_chart:
        genter_quality_data = CategoryChartData()
        genter_quality_data.categories = ['Treasury', 'A', 'BBB']
        genter_quality_data.add_series('Recommended Portfolio', [50.0, 20.0, 30.0])
        chart_s8_rec.chart.replace_data(genter_quality_data)

    # ========================================================
    # SLIDE 9: Hide Municipal State Breakdown Slide
    # The PDF proposal drops State Breakdown completely.
    # Setting show="0" hides it natively without XML package collisions.
    # ========================================================
    print("Hiding Municipal State Breakdown slide...")
    slide9_state = prs.slides[8]
    slide9_state.element.set('show', '0')

    # ========================================================
    # SLIDE 10: Fixed Income Characteristics Comparison
    # ========================================================
    print("Populating Slide 10 (Characteristics Comparison)...")
    slide10 = prs.slides[9]

    def weighted_avg(col_name):
        if col_name not in df.columns:
            return 0.0
        valid = df[col_name].notna() & df['market_value'].notna()
        if not valid.any():
            return 0.0
        return (df.loc[valid, col_name] * df.loc[valid, 'market_value']).sum() / df.loc[valid, 'market_value'].sum()

    avg_ytw = weighted_avg('yield_to_worst')
    avg_ytm = weighted_avg('yield_to_maturity')
    avg_coupon = weighted_avg('coupon_clean')
    avg_effMat = weighted_avg('years_to_eff_maturity')
    avg_effDur = weighted_avg('effective_duration')
    avg_conv = weighted_avg('convexity')

    # Client Values
    c_coupon = get_shape(slide10, "ClientAVGCoupon")
    if c_coupon: update_text_preserve_format(c_coupon, f"{avg_coupon:.2f}%")

    c_ytm = get_shape(slide10, "ClientYTM")
    if c_ytm: update_text_preserve_format(c_ytm, f"{avg_ytm:.2f}%")

    c_ytw = get_shape(slide10, "ClientYTW")
    if c_ytw: update_text_preserve_format(c_ytw, f"{avg_ytw:.2f}%")

    c_tey = get_shape(slide10, "ClientTEY")
    if c_tey: update_text_preserve_format(c_tey, f"{avg_ytw:.2f}%")  # 0.0% tax -> TEY == YTW

    c_effMat = get_shape(slide10, "ClientEffMat")
    if c_effMat: update_text_preserve_format(c_effMat, f"{avg_effMat:.2f} Years")

    c_effDur = get_shape(slide10, "ClientEffDur")
    if c_effDur: update_text_preserve_format(c_effDur, f"{avg_effDur:.2f} Years")

    c_conv = get_shape(slide10, "ClientConv")
    if c_conv: update_text_preserve_format(c_conv, f"{avg_conv:.2f}")

    # Genter Values (Exact targets from PDF Page 9)
    # Average Coupon: 4.42% | Yield to Call: 4.61% | Yrs to Mty: 6.10 | Eff Mty: 6.10 | Eff Dur: 3.98 | Convexity: 0.23
    g_coupon = get_shape(slide10, "GenterAVGCoupon")
    if g_coupon: update_text_preserve_format(g_coupon, "4.42%")

    g_ytm = get_shape(slide10, "GenterYTM")
    if g_ytm: update_text_preserve_format(g_ytm, "4.61%")

    g_ytw = get_shape(slide10, "GenterYTW")
    if g_ytw: update_text_preserve_format(g_ytw, "4.61%")

    g_tey = get_shape(slide10, "GenterTEY")
    if g_tey: update_text_preserve_format(g_tey, "4.61%")

    g_effMat = get_shape(slide10, "GenterEffMat")
    if g_effMat: update_text_preserve_format(g_effMat, "6.10 Years")

    g_effDur = get_shape(slide10, "GenterEffDur")
    if g_effDur: update_text_preserve_format(g_effDur, "3.98 Years")

    g_conv = get_shape(slide10, "GenterConv")
    if g_conv: update_text_preserve_format(g_conv, "0.23")

    # Footnote matching PDF Page 9: 0.0% Federal Tax
    footnote = get_shape(slide10, "TextBox 77")
    if footnote:
        update_text_preserve_format(
            footnote,
            "* Calculated subject to the 0.0% Federal Tax. This calculation does not include any state or local taxes that may apply."
        )

    # Average Quality
    c_qual = get_shape(slide10, "ClientQual")
    if c_qual:
        df_rated = df[df['syn_rating'].isin(RATING_SCORES.keys())].copy()
        df_rated['rating_score'] = df_rated['syn_rating'].map(RATING_SCORES)
        total_rated = df_rated['market_value'].sum()
        if total_rated > 0:
            avg_score = (df_rated['rating_score'] * df_rated['market_value']).sum() / total_rated
            avg_rating = SCORE_TO_RATING.get(round(avg_score), "NR")
        else:
            avg_rating = "NR"
        update_text_preserve_format(c_qual, avg_rating)

    # ========================================================
    # SLIDE 11+: Current Portfolio Investment Detail by Maturity
    # Uses template's Slide 11 for chunk 0, adds slides only if needed
    # ========================================================
    print("Populating Holdings Detail Slides...")
    df_sorted = df.sort_values(by='maturity_date', na_position='last')
    total_mv = df['market_value'].sum()
    chunk_size = 18

    headers = [
        "Quantity", "Symbol / Cusip", "Security Description", "Price", "Market Value",
        "Coupon", "Maturity", "Next Call Date", "YTW", "Ratings", "Modified Duration",
        "Effective Duration", "Convexity", "% of Portfolio"
    ]
    col_widths = [
        Inches(0.65), Inches(1.0), Inches(3.3), Inches(0.6), Inches(0.65),
        Inches(0.7),  Inches(0.9), Inches(0.9), Inches(0.6), Inches(0.6),
        Inches(0.7),  Inches(0.8), Inches(0.8), Inches(0.7)
    ]

    holdings_layout = None
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == "HoldingsDetail":
                holdings_layout = layout
                break
        if holdings_layout:
            break
    if holdings_layout is None:
        holdings_layout = prs.slide_masters[-1].slide_layouts[-1]

    num_chunks = max(1, (len(df_sorted) + chunk_size - 1) // chunk_size)
    for chunk_idx in range(num_chunks):
        chunk = df_sorted.iloc[chunk_idx * chunk_size:(chunk_idx + 1) * chunk_size]

        if chunk_idx == 0 and len(prs.slides) > 10:
            # Use existing Slide 11 from the template
            target_slide = prs.slides[10]
        else:
            target_slide = prs.slides.add_slide(holdings_layout)

        if target_slide.shapes.title:
            target_slide.shapes.title.text = "Current Portfolio Investment Detail by Maturity"

        rows = len(chunk) + 1
        cols = 14

        table_shape = target_slide.shapes.add_table(
            rows, cols, Inches(0.2), Inches(0.9), Inches(12.9), Inches(6.1)
        )
        table = table_shape.table

        for c_idx, width in enumerate(col_widths):
            table.columns[c_idx].width = width

        # Header Row
        for c_idx, header in enumerate(headers):
            cell = table.cell(0, c_idx)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x59, 0x00, 0x2E)

            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for row_idx, (_, bond) in enumerate(chunk.iterrows(), start=1):
            face = f"{bond.get('face_value', 0):,.0f}"
            cusip = str(bond.get('cusip', ''))
            name = str(bond.get('security_name', ''))[:28]
            price = f"${bond.get('current_price', 0):.2f}"
            mv = f"${bond.get('market_value', 0):,.0f}"
            cpn = f"{bond.get('coupon_clean', 0):.3f}%"

            mat_dt = bond.get('maturity_date')
            mat = mat_dt.strftime('%m/%d/%Y') if pd.notna(mat_dt) else "N/A"

            call_dt = bond.get('next_call_date')
            if pd.notna(call_dt) and str(call_dt).strip() not in ["NC", "nan", "NaT"]:
                call = call_dt.strftime('%m/%d/%Y') if hasattr(call_dt, 'strftime') else str(call_dt)[:10]
            else:
                call = "NC"

            ytw = f"{bond.get('yield_to_worst', 0):.2f}%"
            rating = str(bond.get('syn_rating', 'NR'))
            mod_dur = f"{bond.get('modified_duration', 0):.2f}" if pd.notna(bond.get('modified_duration')) else "N/A"
            eff_dur = f"{bond.get('effective_duration', 0):.2f}" if pd.notna(bond.get('effective_duration')) else "N/A"
            conv = f"{bond.get('convexity', 0):.2f}" if pd.notna(bond.get('convexity')) else "N/A"
            pct = f"{(bond.get('market_value', 0) / total_mv) * 100:.1f}%" if total_mv > 0 else "0.0%"

            row_data = [face, cusip, name, price, mv, cpn, mat, call, ytw, rating, mod_dur, eff_dur, conv, pct]

            for c_idx, val in enumerate(row_data):
                cell = table.cell(row_idx, c_idx)
                cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8)

    # ========================================================
    # SAVE PRESENTATION
    # ========================================================
    primary_output = "PAR_Report_Output_Taxable.pptx"
    try:
        prs.save(primary_output)
        print(f"\nSUCCESS: Taxable presentation saved to: {primary_output}")
        return primary_output
    except PermissionError:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_output = f"PAR_Report_Output_Taxable_{ts}.pptx"
        prs.save(fallback_output)
        print(f"\nNOTICE: '{primary_output}' was open in PowerPoint. Saved instead to: {fallback_output}")
        return fallback_output


if __name__ == "__main__":
    generate_taxable_presentation(
        advisor_name="Jeff Kuchinsky, CFP, MBA\nJMK Wealth Advisors",
        client_name="Prospective Client"
    )
